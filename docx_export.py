"""把转录结果排成美观的 .docx：自动分段、时间标签、页脚页码、中西文双字体。

只依赖 python-docx。mode="blocks" 按停顿聚合成自然段（适合阅读/打印），
mode="lines" 一句一行带时间戳（适合逐句精学）。
"""

from __future__ import annotations

import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import env

INK = RGBColor(0x25, 0x24, 0x22)      # 正文
GREY = RGBColor(0x8A, 0x85, 0x7B)     # 次要信息
FAINT = RGBColor(0xB8, 0xB2, 0xA6)    # 时间标签
LATIN = "DejaVu Serif" if sys.platform.startswith("linux") else "Georgia"
CJK = env.cjk_font()                  # 中文按平台取实际存在的字体
SENT_END = tuple("。！？.!?…\"”')')") + (".", "!", "?")


def _set_fonts(run, size, color, bold=False, italic=False, latin=LATIN, cjk=CJK):
    """同时设西文与 eastAsia 字体，否则中文会退回默认。"""
    f = run.font
    f.name = latin
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), cjk)


def _bottom_rule(par, color="D9D3C7", size=6):
    """给段落加一条细分隔线（比插入表格轻）。"""
    p_pr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    # w:pBdr 必须排在 pStyle 之后、spacing 之前
    ref = p_pr.find(qn("w:spacing"))
    if ref is not None:
        ref.addprevious(borders)
    else:
        p_pr.append(borders)


def _page_number(par):
    """页脚插 PAGE 域。"""
    run = par.add_run()
    _set_fonts(run, 8, GREY)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def hms(sec: float) -> str:
    m, s = divmod(int(sec), 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def tidy(text: str) -> str:
    """轻度美化：压空格、首字母大写、结尾补句号。不改词。"""
    t = re.sub(r"\s+", " ", text).strip()
    if not t:
        return ""
    if t[0].isalpha():
        t = t[0].upper() + t[1:]
    if not t.endswith(SENT_END):
        t += "."
    return t


def group(segments, max_gap=2.6, max_span=62.0, max_items=7):
    """按停顿/时长/句数把碎句聚合成块，返回 [(start, end, text), ...]。"""
    blocks = []
    cur = []
    for s in segments:
        if not s.text:
            continue
        if cur and (s.start - cur[-1].end > max_gap
                    or s.end - cur[0].start > max_span
                    or len(cur) >= max_items):
            blocks.append(cur)
            cur = []
        cur.append(s)
    if cur:
        blocks.append(cur)
    return [(b[0].start, b[-1].end, " ".join(x.text.strip() for x in b)) for b in blocks]


def pretty_title(name: str) -> str:
    """把 record-2026_10_18 09_45_39.webm 这类机器名变成可读标题。"""
    stem = re.sub(r"\.[A-Za-z0-9]{2,5}$", "", name)
    m = re.search(r"(\d{4})[-_](\d{2})[-_](\d{2})[ _]?(\d{2})[._:]?(\d{2})", stem)
    if m:
        return f"录音 {m.group(1)}-{m.group(2)}-{m.group(3)} {m.group(4)}:{m.group(5)}"
    return stem


def short_engine(engine: str) -> str:
    """引擎串只留模型名，去掉耗时/加载这些技术尾巴。"""
    e = (engine or "").split("·")[0].strip().replace("本地", "")
    if e and "whisper" not in e.lower():
        e = f"whisper {e}"
    return e or "whisper"


LANGS = {"es": "西班牙语", "zh": "中文", "en": "英语", "ja": "日语", "ko": "韩语",
         "fr": "法语", "de": "德语", "pt": "葡萄牙语", "ru": "俄语", "it": "意大利语",
         "ar": "阿拉伯语"}


def lang_name(code: str) -> str:
    c = (code or "").lower()
    return f"{LANGS.get(c, c or '未知')} ({c.upper()})" if c else "未知"


def build(result, mode: str = "blocks") -> Document:
    """Result -> Document。"""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    # 页脚页码
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_number(fp)

    # 标题
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run(pretty_title(result.source.name))
    _set_fonts(r, 17, INK, bold=True, latin=LATIN, cjk=CJK)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    _bottom_rule(meta)
    r = meta.add_run("  ·  ".join([
        f"语言 {lang_name(result.language)}",
        f"时长 {hms(result.duration)}",
        f"{len(result.segments)} 句",
        short_engine(result.engine),
    ]))
    _set_fonts(r, 9, GREY)

    if mode == "lines":
        for s in result.segments:
            if not s.text:
                continue
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(3)
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.25
            r = p.add_run(f"{hms(s.start)}  ")
            _set_fonts(r, 8.5, FAINT)
            r = p.add_run(tidy(s.text))
            _set_fonts(r, 10.5, INK)
        return doc

    for start, end, text in group(result.segments):
        lab = doc.add_paragraph()
        lab.paragraph_format.space_before = Pt(9)
        lab.paragraph_format.space_after = Pt(1)
        r = lab.add_run(f"{hms(start)} – {hms(end)}")
        _set_fonts(r, 8.5, FAINT, bold=False)

        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(7)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.34
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(tidy(text))
        _set_fonts(r, 11, INK)
    return doc


def save(result, path, mode: str = "blocks"):
    build(result, mode).save(str(path))
    return path
